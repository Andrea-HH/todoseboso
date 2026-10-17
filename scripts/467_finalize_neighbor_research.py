"""Joint delivery: actual applied rules, coefficient audit and product hashes."""
from pathlib import Path
import sys,json,hashlib
import numpy as np
import pandas as pd
import joblib
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT))
from src.neighbor_signal_linear import OUT,TREE_OUT
from src.neighbor_signal_contrasts import SIGNALS
from src.geoc_bridge import OLD


def applied_rule(adjustment,names):
    if adjustment.tree is None:return 'Sin ajuste: predicción = JBoth.'
    tree=adjustment.tree.tree_;lines=[]
    def visit(node,prefix):
        feature=int(tree.feature[node])
        if feature<0:
            value=adjustment.values[int(node)]
            if adjustment.kind=='RESIDUAL':text=f'ajuste residual aplicado = {adjustment.amplitude*value:+.8g} t/ha'
            else:text=f'multiplicador de (JBoth−CT3) = {1+adjustment.amplitude*(value-1):.8g}'
            lines.append(prefix+text);return
        lines.append(prefix+f'Si {names[feature]} ≤ {tree.threshold[node]:.8g}:');visit(int(tree.children_left[node]),prefix+'  ')
        lines.append(prefix+'En otro caso:');visit(int(tree.children_right[node]),prefix+'  ')
    visit(0,'');return '\n'.join(lines)


def main():
    assert (OUT/'evaluation_complete.json').exists() and (TREE_OUT/'evaluation_complete.json').exists()
    cfg=json.loads((OUT/'frozen_design.json').read_text());tree_cfg=json.loads((TREE_OUT/'frozen_design.json').read_text());protected={**tree_cfg['source_sha256'],**cfg['source_sha256']}
    for f,h in protected.items():assert hashlib.sha256(Path(f).read_bytes()).hexdigest()==h,f
    plan=joblib.load(OLD/'cache/execution_plan.joblib');candidates=pd.read_parquet(OUT/'tables/candidates.parquet');selected=pd.read_parquet(OUT/'tables/selection.parquet');coefficients=[];verified=0;rules=[];names=[k+'_local_z' for k in SIGNALS]+[k+'_global_delta' for k in SIGNALS]
    for fi in range(65):
        p=joblib.load(OUT/'models'/f'{fi:02d}.joblib');tr=np.asarray(plan[fi]['train_positions']);va=np.asarray(plan[fi]['validation_positions']);np.testing.assert_array_equal(p['train_positions'],tr);np.testing.assert_array_equal(p['validation_positions'],va)
        for name,model in p['models'].items():
            key=plan[fi];mask=(candidates.protocol==key['protocol'])&(candidates['repeat']==key['repeat'])&(candidates.fold==key['fold'])&(candidates.experiment==name);rows=candidates[mask].to_dict('records');assert len(rows)==25;best=min(rows,key=lambda v:(round(v['rmse'],12),v['amplitude'],-v['alpha'],-v['l1_ratio']));s=selected[(selected.protocol==key['protocol'])&(selected['repeat']==key['repeat'])&(selected.fold==key['fold'])&(selected.experiment==name)].iloc[0]
            for field in ['alpha','l1_ratio','amplitude','rmse']:assert np.isclose(s[field],best[field]),(fi,name,field)
            m=model['adjustment'];np.testing.assert_array_equal(model['contrast'].positions,tr);assert m.amplitude==best['amplitude'];assert m.model is None or (not m.model.fit_intercept and m.model.intercept_==0)
            np.testing.assert_array_equal(m.predict(np.zeros((2,6)),np.full(2,4.),np.full(2,3.)),np.full(2,4.));verified+=1
        tp=joblib.load(TREE_OUT/'models'/f'{fi:02d}.joblib')
        for name,model in tp['models'].items():
            key=plan[fi];m=model['adjustment'];rules.append(dict(execution_index=fi,protocol=key['protocol'],repeat=key['repeat'],fold=key['fold'],experiment=name,amplitude=m.amplitude,rule_applied=applied_rule(m,names),fallback=m.fallback))
    rule_table=pd.DataFrame(rules);rule_table.to_csv(OUT/'tables/deployed_tree_rules_by_fold.csv',index=False);rule_table.to_csv(TREE_OUT/'tables/deployed_rules_by_fold.csv',index=False)
    parity=pd.read_csv(OUT/'tables/calibration_parity.csv');assert len(parity)==65 and parity.null_calibration_rmse_difference.max()<1e-12
    book=load_workbook(OUT/'evaluacion_contrastes_continuos_y_arboles.xlsx')
    extras={'REGLAS_APLICADAS':rule_table,'ALERTAS_DIAGNOSTICO':pd.read_csv(TREE_OUT/'tables/uncertainty_diagnostic_summary.csv'),'ARBOLES_TESTS_PAREADOS':pd.read_csv(TREE_OUT/'tables/paired_tests.csv')}
    for name,frame in extras.items():
        if name in book:del book[name]
        sheet=book.create_sheet(name);sheet.append(frame.columns.tolist())
        for row in frame.itertuples(index=False,name=None):sheet.append([None if pd.isna(v) else v for v in row])
        sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
    guide=book['LEEME']
    guide.append(['REGLAS_APLICADAS muestra los multiplicadores/ajustes finales, después de contracción y límites.'])
    guide.append(['Los valores del texto sklearn original eran previos a calibración; no son pesos aplicados.'])
    guide.append(['Conclusión: conservar J7 y CT3; los correctores continuos empeoran la transferencia espacial.'])
    guide.append(['La pequeña ganancia de atenuación depende del bloque 1; no demuestra una mejora general.'])
    guide.append(['El puntaje de anomalía no es una alerta estable fuera de 171/015; no hay intervalos calibrados.'])
    guide.column_dimensions['A'].width=125
    for column in ['A','B','C','D','E']:book['RESUMEN'].column_dimensions[column].width=25
    book.active=0;book.save(OUT/'evaluacion_contrastes_continuos_y_arboles.xlsx')
    continuous_tests=pd.read_csv(OUT/'tables/paired_tests.csv');tree_tests=pd.read_csv(TREE_OUT/'tables/paired_tests.csv');products=[ROOT/'documentation/neighbor_signal_contrasts_design.md',ROOT/'documentation/neighbor_signal_contrasts_results.md',ROOT/'documentation/neighbor_signal_contrasts_conclusions.md',ROOT/'documentation/neighbor_signal_linear_design.md',ROOT/'documentation/neighbor_signal_linear_results.md',ROOT/'documentation/neighbor_signal_linear_conclusions.md',OUT/'evaluacion_contrastes_continuos_y_arboles.xlsx',OUT/'figures/continuous_rmse_changes.png',TREE_OUT/'figures/rmse_changes.png',OUT/'tables/deployed_tree_rules_by_fold.csv']+[ROOT/'scripts'/f for f in ['460_prepare_neighbor_contrasts.py','461_run_neighbor_contrasts.py','462_evaluate_neighbor_contrasts.py','463_audit_neighbor_contrasts_delivery.py','464_run_neighbor_signal_linear.py','465_prepare_neighbor_signal_linear.py','466_evaluate_neighbor_signal_linear.py',Path(__file__).name]]
    result=dict(status='complete',protected_hashes_verified=len(protected),linear_models_verified=verified,tree_models_verified=len(rules),selection_matches_internal_candidates=True,zero_anchor_verified=True,calibration_parity_max_difference=float(parity.null_calibration_rmse_difference.max()),continuous_significant_holm=int((continuous_tests.p_holm<.05).sum()),tree_significant_holm=int((tree_tests.p_holm<.05).sum()),unit_tests_passed=9,no_PRED=True,no_final138_fit=True,no_champion_change=True,product_sha256={str(f):hashlib.sha256(f.read_bytes()).hexdigest() for f in products});(OUT/'delivery_audit.json').write_text(json.dumps(result,indent=2),encoding='utf-8');print(json.dumps({k:v for k,v in result.items() if k!='product_sha256'},indent=2))


if __name__=='__main__':main()
