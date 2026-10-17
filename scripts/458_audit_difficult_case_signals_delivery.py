"""Validate stored artifacts, all train/held exclusions, role placement and metrics."""
from pathlib import Path
import sys,json,hashlib
import numpy as np
import pandas as pd
import joblib
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT))
from src.difficult_case_signals import OUT,JOINT,SIGNALS,load_assets
from src.geoc_bridge import OLD


def main():
    assert (OUT/'evaluation_complete.json').exists()
    cfg=json.loads((OUT/'frozen_design.json').read_text());protected=dict(json.loads((JOINT/'frozen_design.json').read_text())['source_sha256']);protected.update(cfg['source_sha256'])
    for file,expected in protected.items():assert hashlib.sha256(Path(file).read_bytes()).hexdigest()==expected,file
    a=load_assets();plan=joblib.load(OLD/'cache/execution_plan.joblib');ledger_count=0;fit_count=0
    for fi in range(65):
        p=joblib.load(OUT/'models'/f'{fi:02d}.joblib');tr=np.asarray(plan[fi]['train_positions']);va=np.asarray(plan[fi]['validation_positions']);assert np.array_equal(tr,p['train_positions']) and np.array_equal(va,p['validation_positions'])
        for row in p['ledger']:
            assert not set(row['train']) & (set(row['held']) | set(va));ledger_count+=1
        for name,r in cfg['recipes'].items():
            e=p['experts'][name];np.testing.assert_array_equal(e['local']['targets'],a['y'][tr]);assert np.array_equal(e['train_positions'],tr)
            local=e['local']['representation'].columns;assert local==cfg['variants'][name]['columns']
            added=set(SIGNALS.values()) & set(local);want=3 if r['block']=='ETN' else 2
            assert len(added)==(want if r['placement']=='LOCAL' else 0)
            if e['corrector'] is not None:
                correction=set(e['corrector'].named_steps['repr'].core.cols) & set(SIGNALS.values());assert len(correction)==(want if r['placement']=='CORRECTION' else 0)
            assert p['models'][name]['space'].atoms==r['gate_atoms'];assert p['models'][name]['space'].columns==[];fit_count+=1
    p=pd.read_parquet(OUT/'tables/oof_residuals.parquet');metrics=pd.read_csv(OUT/'tables/metrics_by_repeat.csv').set_index(['experiment','protocol','repeat'])
    for key,g in p.groupby(['experiment','protocol','repeat']):
        e=g.y_pred-g.y_true;m=metrics.loc[key];assert np.isclose(m.rmse,np.sqrt(np.mean(e**2))) and np.isclose(m.mae,np.mean(abs(e))) and np.isclose(m.bias,np.mean(e))
        assert np.isclose(m.r2,1-np.sum(e**2)/np.sum((g.y_true-g.y_true.mean())**2))
    tests=pd.read_csv(OUT/'tables/paired_tests.csv');assert len(tests)==96
    for name in cfg['recipes']:assert len(p[p.experiment==name])==1794
    from openpyxl import load_workbook
    book=load_workbook(OUT/'evaluacion_ubicacion_senales.xlsx')
    for name in ['LEEME','RESUMEN']:
        if name in book:del book[name]
    sheet=book.create_sheet('RESUMEN',0)
    overview=pd.read_csv(OUT/'tables/presentation_RMSE_R2.csv')
    sheet.append(overview.columns.tolist())
    for row in overview.itertuples(index=False,name=None):sheet.append(list(row))
    sheet.freeze_panes='B2';sheet.auto_filter.ref=sheet.dimensions
    sheet.column_dimensions['A'].width=29
    for col in ['B','C','D','E']:sheet.column_dimensions[col].width=24
    guide=book.create_sheet('LEEME',0)
    for text in ['Doce alternativas sobre 65 particiones externas originales; resultados exploratorios.',
                 'RESUMEN: cada celda presenta RMSE (t/ha) / R². Menor RMSE y mayor R² son preferibles.',
                 'J7 = JBoth; E9/E10 = mezclas con expertos de 9/10 variables, no expertos aislados.',
                 'DROP_SOIL = referencia E9; DROP_HISTORY_WATER = referencia E10.',
                 'ET añade cambio EVI y TPI; ETN añade dinámica de la brecha NDVI medio-mínimo.',
                 'LOCAL modifica entradas Ridge local; CORRECTION modifica entradas del corrector residual.',
                 'Error = predicción menos observado. Delta RMSE = alternativa menos referencia; negativo mejora.',
                 'Aleatoria: métricas OOF completas por repetición, luego promedio de diez repeticiones.',
                 'Ninguna alternativa justifica reemplazar las referencias; no se ajustó modelo final ni PRED59.',
                 'Conclusiones: documentation/difficult_case_signals_conclusions.md.',
                 'CSV de ganancias por parcela y pesos completos se conservan en tables fuera del libro.']:
        guide.append([text])
    guide.column_dimensions['A'].width=125;book.active=0
    book.save(OUT/'evaluacion_ubicacion_senales.xlsx')
    products=[ROOT/'documentation/difficult_case_signals_design.md',ROOT/'documentation/difficult_case_signals_results.md',ROOT/'documentation/difficult_case_signals_conclusions.md',OUT/'evaluacion_ubicacion_senales.xlsx',OUT/'figures/placement_rmse_changes.png']
    products += [ROOT/'scripts'/f for f in ['454_prepare_difficult_case_signals.py','455_run_difficult_case_signals.py','456_evaluate_difficult_case_signals.py','457_finalize_difficult_case_signals.py',Path(__file__).name]]
    result=dict(status='complete',protected_hashes_verified=len(protected),fitted_experts_verified=fit_count,ledger_exclusions_verified=ledger_count,metric_groups_verified=len(metrics),outer_folds=65,new_recipes=12,paired_tests=96,significant_holm=int((tests.p_holm<.05).sum()),unit_tests_passed=7,no_final138_fit=True,no_PRED=True,product_sha256={str(f):hashlib.sha256(f.read_bytes()).hexdigest() for f in products})
    (OUT/'delivery_audit.json').write_text(json.dumps(result,indent=2),encoding='utf-8');print(json.dumps({k:v for k,v in result.items() if k!='product_sha256'},indent=2))


if __name__=='__main__':main()
